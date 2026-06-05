"""RAG v12 (Visual+Text RAG) 전체 40문항 테스트 + 채점 + 7개 시스템 비교 + Excel/MD 생성.

실행 순서:
1. 대안천교 40문항 v12 실행 (v10r 검색 + 이미지 + OCR 텍스트 → Gemini 답변)
2. 대안천교 채점
3. 송정교 40문항 v12 실행
4. 송정교 채점
5. v1, v5, v10r, v11, v12, LLM, Proposed_improved 비교표 출력
6. Wilcoxon signed-rank test
7. results/final_comparison_v12.xlsx 저장
8. results/paper_data_summary_v6.md 저장

주의:
- 이 스크립트는 실행 시 실제 API를 호출한다.
- v10r의 검색 파이프라인(청크 캐시, 임베딩)을 재사용한다.
- 답변 생성은 Gemini 3.1 Pro에 PDF 페이지 이미지 + OCR 텍스트를 전달한다.
"""

from __future__ import annotations

import json
import sys
import time
from datetime import datetime
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from scipy import stats

sys.path.insert(0, str(Path(__file__).parent))

import config
from evaluation.judge import judge_all_results
from questions import CATEGORIES, QUESTIONS
from system_a_rag_v10.chunker import build_smart_chunks
from system_a_rag_v10.embeddings import GeminiEmbeddings
from system_a_rag_v10.retriever_v10 import ChapterHybridRetriever
from system_a_rag_v10.upstage_loader import load_upstage_parse_pages
from system_a_rag_v11.retriever_v12 import rag_answer_v12


BRIDGES = ["대안천교", "송정교"]
METRICS = ["accuracy", "completeness", "faithfulness"]
Q_IDS = [q["id"] for q in QUESTIONS]
Q_CAT = {q["id"]: q["category"] for q in QUESTIONS}

ALL_SYSTEMS = {
    "v1": "rag",
    "v5": "rag_v5",
    "v10r": "rag_v10r",
    "v11": "rag_v11",
    "v12": "rag_v12",
    "LLM": "llm",
    "Proposed_improved": "proposed_improved",
}

WILCOXON_PAIRS = [
    ("v1", "v5"),
    ("v1", "v10r"),
    ("v1", "v12"),
    ("v5", "v10r"),
    ("v5", "v12"),
    ("v10r", "v11"),
    ("v10r", "v12"),
    ("v11", "v12"),
    ("v10r", "LLM"),
    ("v12", "LLM"),
    ("v12", "Proposed_improved"),
    ("LLM", "Proposed_improved"),
]


# ── 유틸리티 ──

def _load_json(path: Path, default=None):
    if not path.exists():
        return default
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def _save_json(path: Path, payload: dict | list) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)


def _serialize_docs(docs) -> list[dict]:
    return [{"page_content": doc.page_content, "metadata": doc.metadata} for doc in docs]


def _save_v12_results(bridge_name: str, answers: list[dict], output_path: Path) -> dict:
    payload = {
        "bridge": bridge_name,
        "system": "rag_v12",
        "model": "gemini-3.1-pro-preview",
        "timestamp": datetime.now().isoformat(),
        "description": "v10r retrieval + PDF page images + OCR text → Gemini 3.1 Pro",
        "total_questions": len(answers),
        "results": answers,
    }
    _save_json(output_path, payload)
    return payload


# ── 점수 로드 ──

def load_results(bridge: str, system: str) -> list[dict]:
    path = Path(config.RESULTS_DIR) / f"{bridge}_{system}_results.json"
    data = _load_json(path, default={})
    if isinstance(data, dict):
        return data.get("results", [])
    return data or []


def load_judged(bridge: str, system: str) -> list[dict]:
    path = Path(config.RESULTS_DIR) / f"{bridge}_{system}_judged.json"
    data = _load_json(path, default=[])
    return data or []


def load_scores_map(bridge: str, system: str) -> dict[str, dict]:
    score_map: dict[str, dict] = {}
    for item in load_judged(bridge, system):
        qid = item.get("question_id", "")
        score_map[qid] = {
            "accuracy": item.get("accuracy"),
            "completeness": item.get("completeness"),
            "faithfulness": item.get("faithfulness"),
        }
    return score_map


def avg_scores(scores_map: dict, qids: list[str], metric: str):
    values = [scores_map.get(qid, {}).get(metric) for qid in qids]
    valid = [v for v in values if isinstance(v, (int, float))]
    return (sum(valid) / len(valid)) if valid else None


def overall_avg(scores_map: dict, qids: list[str]):
    values = []
    for metric in METRICS:
        values.extend(
            v for v in (scores_map.get(qid, {}).get(metric) for qid in qids)
            if isinstance(v, (int, float))
        )
    return (sum(values) / len(values)) if values else None


def wilcoxon_test(scores_a: list, scores_b: list):
    paired = [
        (a, b) for a, b in zip(scores_a, scores_b)
        if isinstance(a, (int, float)) and isinstance(b, (int, float))
    ]
    if len(paired) < 5:
        return None, None
    a = np.array([p[0] for p in paired], dtype=float)
    b = np.array([p[1] for p in paired], dtype=float)
    diffs = a - b
    diffs = diffs[diffs != 0]
    if len(diffs) < 5:
        return None, None
    try:
        stat, p = stats.wilcoxon(diffs)
        return stat, p
    except Exception:
        return None, None


def estimate_cost(bridge: str, system_name: str) -> dict:
    results = load_results(bridge, system_name)
    total_input = sum(item.get("input_tokens", 0) for item in results)
    total_output = sum(item.get("output_tokens", 0) for item in results)
    elapsed_vals = [
        item.get("elapsed_sec", 0) for item in results
        if isinstance(item.get("elapsed_sec", 0), (int, float))
    ]
    avg_elapsed = (sum(elapsed_vals) / len(elapsed_vals)) if elapsed_vals else None

    if system_name in {"rag", "rag_v3", "rag_v5", "rag_v10", "rag_v10r"}:
        model_name = "gpt-5.4"
    elif system_name in {"rag_v11", "rag_v12"}:
        model_name = "gemini-3.1-pro-preview"
    else:
        model_name = "gemini-3.1-pro-preview"

    costs = config.MODEL_COSTS.get(model_name, {"input": 0, "output": 0})
    input_cost = total_input / 1_000_000 * costs["input"]
    output_cost = total_output / 1_000_000 * costs["output"]

    return {
        "model": model_name,
        "input_tokens": total_input,
        "output_tokens": total_output,
        "avg_elapsed_sec": avg_elapsed,
        "total_cost": input_cost + output_cost,
    }


def analyze_v12_retrieval(bridge: str) -> list[dict]:
    rows = []
    for item in load_results(bridge, "rag_v12"):
        qid = item.get("question_id", "")
        chapters = item.get("predicted_chapters", []) or []
        pages = item.get("retrieved_pages", []) or []
        image_pages = item.get("image_pages", []) or []
        table_chunks = item.get("used_table_chunks", []) or []
        rows.append({
            "qid": qid,
            "category": Q_CAT.get(qid, ""),
            "chapters": chapters,
            "n_chapters": len(chapters),
            "retrieved_pages": pages,
            "n_pages": len(pages),
            "image_pages": image_pages,
            "n_image_pages": len(image_pages),
            "n_table_chunks": len(table_chunks),
            "elapsed_sec": item.get("elapsed_sec"),
            "input_tokens": item.get("input_tokens", 0),
            "output_tokens": item.get("output_tokens", 0),
        })
    return rows


def summarize_v12_retrieval(rows: list[dict]) -> dict:
    if not rows:
        return {
            "n_questions": 0, "avg_chapters": None, "avg_pages": None,
            "avg_image_pages": None, "avg_tables": None, "chapter_counts": {},
        }
    chapter_counts: dict[int, int] = {}
    for row in rows:
        for ch in row.get("chapters", []):
            chapter_counts[ch] = chapter_counts.get(ch, 0) + 1
    return {
        "n_questions": len(rows),
        "avg_chapters": sum(r["n_chapters"] for r in rows) / len(rows),
        "avg_pages": sum(r["n_pages"] for r in rows) / len(rows),
        "avg_image_pages": sum(r["n_image_pages"] for r in rows) / len(rows),
        "avg_tables": sum(r["n_table_chunks"] for r in rows) / len(rows),
        "chapter_counts": dict(sorted(chapter_counts.items())),
    }


# ── v12 실행 파이프라인 ──

def run_rag_v12_full_pipeline(bridge_name: str, question_ids: list[str] | None = None) -> dict:
    if not config.GOOGLE_API_KEY:
        raise RuntimeError("GOOGLE_API_KEY is not set")

    from langchain_core.documents import Document

    active_questions = [q for q in QUESTIONS if q["id"] in (question_ids or Q_IDS)]
    results_dir = Path(config.RESULTS_DIR)
    results_dir.mkdir(parents=True, exist_ok=True)
    output_path = results_dir / f"{bridge_name}_rag_v12_results.json"

    existing_payload = _load_json(output_path, default={}) or {}
    existing_results = existing_payload.get("results", []) if isinstance(existing_payload, dict) else []
    answers_by_qid = {item.get("question_id", ""): item for item in existing_results}
    pending_questions = [q for q in active_questions if q["id"] not in answers_by_qid]

    if not pending_questions and len(answers_by_qid) >= len(active_questions):
        print(f"\n⏭️ [v12] {bridge_name} 이미 {len(active_questions)}문항 결과 존재 → 스킵")
        ordered = [answers_by_qid[q["id"]] for q in active_questions if q["id"] in answers_by_qid]
        return _save_v12_results(bridge_name, ordered, output_path)

    print(f"\n🚀 [RAG v12 Full] {bridge_name} 시작 ({len(active_questions)}문항)")
    if answers_by_qid:
        print(f"🔁 이어서 실행: 기존 {len(answers_by_qid)}문항, 남은 {len(pending_questions)}문항")

    # 청크 로드 (v10r 캐시 재사용)
    chunk_cache = Path(config.PROJECT_ROOT) / "data" / bridge_name / "upstage_v10_chunks.json"
    if chunk_cache.exists():
        print("  📦 청크 캐시 로드 중...")
        raw = json.loads(chunk_cache.read_text(encoding="utf-8"))
        docs = [Document(page_content=d["page_content"], metadata=d["metadata"]) for d in raw]
    else:
        print("  📄 Upstage 파싱 + 스마트 청킹...")
        upstage_pages, failed_pages = load_upstage_parse_pages(bridge_name)
        print(f"  Upstage pages: success={len(upstage_pages)}, failed={len(failed_pages)}")
        docs = build_smart_chunks(upstage_pages)
    print(f"  🧩 Smart chunks: {len(docs)}개")

    # 임베딩 로드 (v10r 캐시 재사용)
    emb_cache = Path(config.PROJECT_ROOT) / "data" / bridge_name / "v10_embeddings.npy"
    embedder = GeminiEmbeddings()
    if emb_cache.exists():
        print("  📦 임베딩 캐시 로드 중...")
        vectors = np.load(str(emb_cache))
    else:
        print("  🔢 Gemini 임베딩 생성 중...")
        texts = [doc.page_content for doc in docs]
        vectors = np.array(embedder.embed_documents(texts), dtype="float32")
        norms = np.linalg.norm(vectors, axis=1, keepdims=True)
        norms[norms == 0] = 1.0
        vectors = vectors / norms
        np.save(str(emb_cache), vectors)
    print(f"  🔢 Embeddings: {vectors.shape}")

    retriever = ChapterHybridRetriever(docs, vectors, embedder)

    for idx, question in enumerate(pending_questions, 1):
        print(f"  ▶ {question['id']} ({idx}/{len(pending_questions)})")
        answer_result = None

        for attempt in range(3):
            try:
                answer_result = rag_answer_v12(
                    question=question["text"],
                    retriever=retriever,
                    all_docs=docs,
                    bridge_name=bridge_name,
                    question_id=question["id"],
                )
                answer_result["category"] = question["category"]
                break
            except Exception as exc:
                wait_sec = 30 if "429" in str(exc) else 10
                if attempt < 2:
                    print(f"    ⚠️ 에러 (시도 {attempt + 1}/3): {exc}")
                    time.sleep(wait_sec)
                    continue
                answer_result = {
                    "question_id": question["id"],
                    "question": question["text"],
                    "answer": f"ERROR: {exc}",
                    "category": question["category"],
                    "predicted_chapters": [],
                    "subqueries": [],
                    "retrieved_pages": [],
                    "image_pages": [],
                    "n_image_pages": 0,
                    "total_image_kb": 0,
                    "used_table_chunks": [],
                    "bridge": bridge_name,
                    "system": "rag_v12",
                    "elapsed_sec": 0,
                    "input_tokens": 0,
                    "output_tokens": 0,
                }

        answers_by_qid[question["id"]] = answer_result
        n_img = answer_result.get("n_image_pages", 0)
        img_kb = answer_result.get("total_image_kb", 0)
        print(
            f"    chapters={answer_result.get('predicted_chapters', [])} | "
            f"images={n_img}장({img_kb:.0f}KB) | "
            f"{answer_result.get('elapsed_sec', 0):.1f}s"
        )

        ordered_answers = [answers_by_qid[q["id"]] for q in active_questions if q["id"] in answers_by_qid]
        if len(ordered_answers) % 10 == 0 or idx == len(pending_questions):
            _save_v12_results(bridge_name, ordered_answers, output_path)
            print(f"    💾 중간 저장 ({len(ordered_answers)}/{len(active_questions)})")

        if idx < len(pending_questions):
            time.sleep(3)

    ordered_answers = [answers_by_qid[q["id"]] for q in active_questions if q["id"] in answers_by_qid]
    final_payload = _save_v12_results(bridge_name, ordered_answers, output_path)
    print(f"  ✅ [v12] {bridge_name} 완료: {output_path}")
    return final_payload


# ── Excel 생성 ──

def create_excel(all_data: dict, retrieval_analysis: dict, output_path: str) -> None:
    wb = Workbook()
    header_font = Font(bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    center = Alignment(horizontal="center", vertical="center")
    best_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    best_font = Font(bold=True, color="006100")
    sig_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    def style_header(ws, row: int, cols: int) -> None:
        for col in range(1, cols + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            cell.border = thin_border

    def style_cell(ws, row: int, col: int, is_best: bool = False) -> None:
        cell = ws.cell(row=row, column=col)
        cell.border = thin_border
        cell.alignment = center
        if is_best:
            cell.fill = best_fill
            cell.font = best_font

    system_labels = list(ALL_SYSTEMS.keys())

    # Sheet 1: 전체 성능 요약
    ws1 = wb.active
    ws1.title = "전체 성능 요약"
    for col in range(1, len(system_labels) + 3):
        ws1.column_dimensions[chr(64 + col)].width = 16

    row = 1
    for bridge in BRIDGES:
        ws1.cell(row=row, column=1, value=f"■ {bridge}").font = Font(bold=True, size=13)
        row += 1
        headers = ["메트릭"] + system_labels
        for idx, value in enumerate(headers, 1):
            ws1.cell(row=row, column=idx, value=value)
        style_header(ws1, row, len(headers))
        row += 1

        bridge_data = all_data.get(bridge, {})
        for metric in METRICS:
            metric_kr = {"accuracy": "정확성", "completeness": "완전성", "faithfulness": "충실성"}[metric]
            ws1.cell(row=row, column=1, value=metric_kr).border = thin_border
            values = [avg_scores(bridge_data.get(ALL_SYSTEMS[label], {}), Q_IDS, metric) for label in system_labels]
            best_val = max((v for v in values if v is not None), default=None)
            for idx, value in enumerate(values, 2):
                if value is None:
                    ws1.cell(row=row, column=idx, value="-")
                    style_cell(ws1, row, idx)
                else:
                    ws1.cell(row=row, column=idx, value=round(value, 2))
                    style_cell(ws1, row, idx, value == best_val)
            row += 1

        ws1.cell(row=row, column=1, value="전체 평균").font = Font(bold=True)
        ws1.cell(row=row, column=1).border = thin_border
        values = [overall_avg(bridge_data.get(ALL_SYSTEMS[label], {}), Q_IDS) for label in system_labels]
        best_val = max((v for v in values if v is not None), default=None)
        for idx, value in enumerate(values, 2):
            if value is None:
                ws1.cell(row=row, column=idx, value="-")
                style_cell(ws1, row, idx)
            else:
                ws1.cell(row=row, column=idx, value=round(value, 2))
                style_cell(ws1, row, idx, value == best_val)
        row += 2

    # Sheet 2: 문항별 상세
    ws2 = wb.create_sheet("문항별 상세")
    for col in range(1, len(system_labels) + 4):
        ws2.column_dimensions[chr(64 + col)].width = 14

    row = 1
    for bridge in BRIDGES:
        ws2.cell(row=row, column=1, value=f"■ {bridge}").font = Font(bold=True, size=13)
        row += 1
        headers = ["문항", "카테고리", "메트릭"] + system_labels
        for idx, value in enumerate(headers, 1):
            ws2.cell(row=row, column=idx, value=value)
        style_header(ws2, row, len(headers))
        row += 1

        bridge_data = all_data.get(bridge, {})
        for qid in Q_IDS:
            for metric_idx, metric in enumerate(METRICS):
                metric_kr = {"accuracy": "정확성", "completeness": "완전성", "faithfulness": "충실성"}[metric]
                if metric_idx == 0:
                    ws2.cell(row=row, column=1, value=qid).border = thin_border
                    ws2.cell(row=row, column=2, value=Q_CAT.get(qid, "")).border = thin_border
                else:
                    ws2.cell(row=row, column=1).border = thin_border
                    ws2.cell(row=row, column=2).border = thin_border
                ws2.cell(row=row, column=3, value=metric_kr).border = thin_border
                values = [bridge_data.get(ALL_SYSTEMS[label], {}).get(qid, {}).get(metric) for label in system_labels]
                best_val = max((v for v in values if v is not None), default=None)
                for idx, value in enumerate(values, 4):
                    if value is None:
                        ws2.cell(row=row, column=idx, value="-")
                        style_cell(ws2, row, idx)
                    else:
                        ws2.cell(row=row, column=idx, value=value)
                        style_cell(ws2, row, idx, value == best_val)
                row += 1

    # Sheet 3: Wilcoxon Test
    ws3 = wb.create_sheet("Wilcoxon Test")
    for col in "ABCDEF":
        ws3.column_dimensions[col].width = 18

    row = 1
    for bridge in BRIDGES:
        ws3.cell(row=row, column=1, value=f"■ {bridge}").font = Font(bold=True, size=13)
        row += 1
        headers = ["교량", "시스템 쌍", "메트릭", "통계량", "p-value", "유의미(p<0.05)"]
        for idx, value in enumerate(headers, 1):
            ws3.cell(row=row, column=idx, value=value)
        style_header(ws3, row, len(headers))
        row += 1

        bridge_data = all_data.get(bridge, {})
        for left, right in WILCOXON_PAIRS:
            left_map = bridge_data.get(ALL_SYSTEMS[left], {})
            right_map = bridge_data.get(ALL_SYSTEMS[right], {})
            for metric in METRICS:
                left_scores = [left_map.get(qid, {}).get(metric) for qid in Q_IDS]
                right_scores = [right_map.get(qid, {}).get(metric) for qid in Q_IDS]
                stat, p = wilcoxon_test(left_scores, right_scores)
                metric_kr = {"accuracy": "정확성", "completeness": "완전성", "faithfulness": "충실성"}[metric]
                ws3.cell(row=row, column=1, value=bridge)
                ws3.cell(row=row, column=2, value=f"{left} vs {right}")
                ws3.cell(row=row, column=3, value=metric_kr)
                ws3.cell(row=row, column=4, value=round(stat, 4) if stat is not None else "-")
                ws3.cell(row=row, column=5, value=round(p, 6) if p is not None else "-")
                significant = "Yes" if p is not None and p < 0.05 else ("No" if p is not None else "-")
                ws3.cell(row=row, column=6, value=significant)
                for col in range(1, 7):
                    ws3.cell(row=row, column=col).border = thin_border
                    ws3.cell(row=row, column=col).alignment = center
                    if significant == "Yes":
                        ws3.cell(row=row, column=col).fill = sig_fill
                row += 1
        row += 1

    # Sheet 4: 카테고리별 분석
    ws4 = wb.create_sheet("카테고리별 분석")
    for col in range(1, len(system_labels) + 3):
        ws4.column_dimensions[chr(64 + col)].width = 16

    row = 1
    for bridge in BRIDGES:
        ws4.cell(row=row, column=1, value=f"■ {bridge}").font = Font(bold=True, size=13)
        row += 1
        bridge_data = all_data.get(bridge, {})
        for category in CATEGORIES:
            ws4.cell(row=row, column=1, value=f"▸ {category}").font = Font(bold=True)
            row += 1
            headers = ["메트릭"] + system_labels
            for idx, value in enumerate(headers, 1):
                ws4.cell(row=row, column=idx, value=value)
            style_header(ws4, row, len(headers))
            row += 1
            cat_qids = [q["id"] for q in QUESTIONS if q["category"] == category]
            for metric in METRICS:
                metric_kr = {"accuracy": "정확성", "completeness": "완전성", "faithfulness": "충실성"}[metric]
                ws4.cell(row=row, column=1, value=metric_kr).border = thin_border
                values = [avg_scores(bridge_data.get(ALL_SYSTEMS[label], {}), cat_qids, metric) for label in system_labels]
                best_val = max((v for v in values if v is not None), default=None)
                for idx, value in enumerate(values, 2):
                    if value is None:
                        ws4.cell(row=row, column=idx, value="-")
                        style_cell(ws4, row, idx)
                    else:
                        ws4.cell(row=row, column=idx, value=round(value, 2))
                        style_cell(ws4, row, idx, value == best_val)
                row += 1
            row += 1

    # Sheet 5: 비용 분석
    ws5 = wb.create_sheet("비용 분석")
    widths = {"A": 12, "B": 20, "C": 24, "D": 14, "E": 14, "F": 16, "G": 12}
    for col, width in widths.items():
        ws5.column_dimensions[col].width = width
    headers = ["교량", "시스템", "모델", "입력 토큰", "출력 토큰", "평균 응답시간(초)", "비용(USD)"]
    for idx, value in enumerate(headers, 1):
        ws5.cell(row=1, column=idx, value=value)
    style_header(ws5, 1, len(headers))

    row = 2
    for bridge in BRIDGES:
        for label in system_labels:
            info = estimate_cost(bridge, ALL_SYSTEMS[label])
            ws5.cell(row=row, column=1, value=bridge)
            ws5.cell(row=row, column=2, value=label)
            ws5.cell(row=row, column=3, value=info["model"])
            ws5.cell(row=row, column=4, value=info["input_tokens"])
            ws5.cell(row=row, column=5, value=info["output_tokens"])
            ws5.cell(row=row, column=6, value=round(info["avg_elapsed_sec"], 2) if info["avg_elapsed_sec"] is not None else "-")
            ws5.cell(row=row, column=7, value=round(info["total_cost"], 4))
            for col in range(1, 8):
                ws5.cell(row=row, column=col).border = thin_border
                ws5.cell(row=row, column=col).alignment = center
            row += 1
        row += 1

    # Sheet 6: v12 검색/이미지 분석
    ws6 = wb.create_sheet("v12 검색 분석")
    widths = {"A": 10, "B": 14, "C": 14, "D": 10, "E": 10, "F": 24, "G": 12, "H": 12}
    for col, width in widths.items():
        ws6.column_dimensions[col].width = width

    row = 1
    for bridge in BRIDGES:
        ws6.cell(row=row, column=1, value=f"■ {bridge}").font = Font(bold=True, size=13)
        row += 1
        headers = ["문항", "카테고리", "예측 장", "장 수", "이미지 수", "이미지 페이지", "표 청크 수", "응답시간(초)"]
        for idx, value in enumerate(headers, 1):
            ws6.cell(row=row, column=idx, value=value)
        style_header(ws6, row, len(headers))
        row += 1

        for item in retrieval_analysis.get(bridge, []):
            ws6.cell(row=row, column=1, value=item["qid"])
            ws6.cell(row=row, column=2, value=item["category"])
            ws6.cell(row=row, column=3, value=", ".join(str(ch) for ch in item["chapters"]) or "-")
            ws6.cell(row=row, column=4, value=item["n_chapters"])
            ws6.cell(row=row, column=5, value=item["n_image_pages"])
            ws6.cell(row=row, column=6, value=", ".join(str(p) for p in item["image_pages"]) or "-")
            ws6.cell(row=row, column=7, value=item["n_table_chunks"])
            ws6.cell(row=row, column=8, value=round(item["elapsed_sec"], 2) if isinstance(item["elapsed_sec"], (int, float)) else "-")
            for col in range(1, 9):
                ws6.cell(row=row, column=col).border = thin_border
            row += 1
        row += 2

    wb.save(output_path)
    print(f"\n💾 Excel 저장: {output_path}")


# ── Markdown 생성 ──

def create_markdown(all_data: dict, retrieval_analysis: dict, output_path: str) -> None:
    system_labels = list(ALL_SYSTEMS.keys())
    lines: list[str] = []
    lines.append("# Paper Data Summary v6 — v12 (Visual+Text RAG) 포함 7개 시스템 비교\n")
    lines.append(f"생성: {datetime.now().strftime('%Y-%m-%d %H:%M')}\n")

    # Section 1
    lines.append("## 1. 전체 성능 요약\n")
    for bridge in BRIDGES:
        lines.append(f"### {bridge}\n")
        lines.append("| 메트릭 | " + " | ".join(system_labels) + " |")
        lines.append("| --- | " + " | ".join(["---"] * len(system_labels)) + " |")
        bridge_data = all_data.get(bridge, {})
        for metric in METRICS:
            label = {"accuracy": "정확성", "completeness": "완전성", "faithfulness": "충실성"}[metric]
            values = [avg_scores(bridge_data.get(ALL_SYSTEMS[s], {}), Q_IDS, metric) for s in system_labels]
            best_val = max((v for v in values if v is not None), default=None)
            rendered = [f"**{v:.2f}**" if v == best_val else (f"{v:.2f}" if v is not None else "-") for v in values]
            lines.append(f"| {label} | " + " | ".join(rendered) + " |")
        values = [overall_avg(bridge_data.get(ALL_SYSTEMS[s], {}), Q_IDS) for s in system_labels]
        best_val = max((v for v in values if v is not None), default=None)
        rendered = [f"**{v:.2f}**" if v == best_val else (f"{v:.2f}" if v is not None else "-") for v in values]
        lines.append("| **전체 평균** | " + " | ".join(rendered) + " |")
        lines.append("")

    # Section 2
    lines.append("## 2. Wilcoxon Signed-Rank Test\n")
    for bridge in BRIDGES:
        lines.append(f"### {bridge}\n")
        lines.append("| 시스템 쌍 | 정확성 p | 완전성 p | 충실성 p |")
        lines.append("| --- | --- | --- | --- |")
        bridge_data = all_data.get(bridge, {})
        for left, right in WILCOXON_PAIRS:
            row_cells = [f"{left} vs {right}"]
            for metric in METRICS:
                left_scores = [bridge_data.get(ALL_SYSTEMS[left], {}).get(qid, {}).get(metric) for qid in Q_IDS]
                right_scores = [bridge_data.get(ALL_SYSTEMS[right], {}).get(qid, {}).get(metric) for qid in Q_IDS]
                _, p = wilcoxon_test(left_scores, right_scores)
                if p is None:
                    row_cells.append("-")
                elif p < 0.05:
                    row_cells.append(f"**{p:.4f}**")
                else:
                    row_cells.append(f"{p:.4f}")
            lines.append("| " + " | ".join(row_cells) + " |")
        lines.append("")

    # Section 3
    lines.append("## 3. 카테고리별 전체 평균\n")
    for bridge in BRIDGES:
        lines.append(f"### {bridge}\n")
        lines.append("| 카테고리 | " + " | ".join(system_labels) + " |")
        lines.append("| --- | " + " | ".join(["---"] * len(system_labels)) + " |")
        bridge_data = all_data.get(bridge, {})
        for category in CATEGORIES:
            cat_qids = [q["id"] for q in QUESTIONS if q["category"] == category]
            values = [overall_avg(bridge_data.get(ALL_SYSTEMS[s], {}), cat_qids) for s in system_labels]
            best_val = max((v for v in values if v is not None), default=None)
            rendered = [f"**{v:.2f}**" if v == best_val else (f"{v:.2f}" if v is not None else "-") for v in values]
            lines.append(f"| {category} | " + " | ".join(rendered) + " |")
        lines.append("")

    # Section 4
    lines.append("## 4. v12 검색/이미지 특성 분석\n")
    for bridge in BRIDGES:
        summary = summarize_v12_retrieval(retrieval_analysis.get(bridge, []))
        chapter_hits = ", ".join(f"{ch}장={cnt}" for ch, cnt in summary["chapter_counts"].items()) or "-"
        lines.append(f"### {bridge}\n")
        lines.append(f"- 문항 수: {summary['n_questions']}")
        lines.append(f"- 평균 예측 장 수: {summary['avg_chapters']:.2f}" if summary["avg_chapters"] is not None else "- 평균 예측 장 수: -")
        lines.append(f"- 평균 검색 페이지 수: {summary['avg_pages']:.2f}" if summary["avg_pages"] is not None else "- 평균 검색 페이지 수: -")
        lines.append(f"- 평균 이미지 페이지 수: {summary['avg_image_pages']:.2f}" if summary["avg_image_pages"] is not None else "- 평균 이미지 페이지 수: -")
        lines.append(f"- 평균 표 청크 수: {summary['avg_tables']:.2f}" if summary["avg_tables"] is not None else "- 평균 표 청크 수: -")
        lines.append(f"- 장 히트 분포: {chapter_hits}\n")

    # Section 5
    lines.append("## 5. 비용 및 응답시간\n")
    lines.append("| 교량 | 시스템 | 모델 | 평균 응답시간(초) | 비용(USD) |")
    lines.append("| --- | --- | --- | --- | --- |")
    for bridge in BRIDGES:
        for label in system_labels:
            info = estimate_cost(bridge, ALL_SYSTEMS[label])
            elapsed = f"{info['avg_elapsed_sec']:.2f}" if info["avg_elapsed_sec"] is not None else "-"
            lines.append(f"| {bridge} | {label} | {info['model']} | {elapsed} | ${info['total_cost']:.4f} |")
    lines.append("")

    # Section 6
    lines.append("## 6. 핵심 수치 (Abstract/Conclusion용)\n")
    for bridge in BRIDGES:
        lines.append(f"### {bridge}\n")
        bridge_data = all_data.get(bridge, {})
        for label in system_labels:
            value = overall_avg(bridge_data.get(ALL_SYSTEMS[label], {}), Q_IDS)
            lines.append(f"- {label}: {value:.2f}" if value is not None else f"- {label}: -")
        lines.append("")

    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print(f"💾 Markdown 저장: {output_path}")


# ── main ──

def main() -> None:
    global_start = time.time()
    results_dir = Path(config.RESULTS_DIR)
    results_dir.mkdir(parents=True, exist_ok=True)

    print("=" * 72)
    print("  Full v12 Test: 2 bridges × 40 questions + 7-system comparison")
    print("  v10r retrieval + PDF page images + OCR text → Gemini 3.1 Pro")
    print(f"  시작: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    print("=" * 72)

    # Phase 1: 대안천교 v12 실행
    bridge = "대안천교"
    result_file = results_dir / f"{bridge}_rag_v12_results.json"
    existing = _load_json(result_file, default={}) or {}
    existing_n = existing.get("total_questions", len(existing.get("results", []))) if isinstance(existing, dict) else 0
    if existing_n >= 40:
        print(f"\n⏭️ [v12] {bridge} 이미 {existing_n}문항 결과 존재 → 스킵")
    else:
        print(f"\n{'='*60}")
        print(f"  [실행] v12 — {bridge} (40문항)")
        print(f"{'='*60}")
        t0 = time.time()
        run_rag_v12_full_pipeline(bridge, Q_IDS)
        elapsed = time.time() - t0
        print(f"  ⏱️ 소요: {int(elapsed // 60)}분 {int(elapsed % 60):02d}초")

    # Phase 2: 대안천교 채점
    judged_file = results_dir / f"{bridge}_rag_v12_judged.json"
    judged_existing = _load_json(judged_file, default=[]) or []
    if len(judged_existing) >= 40:
        print(f"\n⏭️ [v12] {bridge} 이미 {len(judged_existing)}문항 채점 완료 → 스킵")
    else:
        print(f"\n{'='*60}")
        print(f"  [채점] v12 — {bridge}")
        print(f"{'='*60}")
        judge_all_results(bridge, str(result_file), "rag_v12")
        time.sleep(5)

    # Phase 3: 송정교 v12 실행
    bridge = "송정교"
    result_file = results_dir / f"{bridge}_rag_v12_results.json"
    existing = _load_json(result_file, default={}) or {}
    existing_n = existing.get("total_questions", len(existing.get("results", []))) if isinstance(existing, dict) else 0
    if existing_n >= 40:
        print(f"\n⏭️ [v12] {bridge} 이미 {existing_n}문항 결과 존재 → 스킵")
    else:
        print(f"\n{'='*60}")
        print(f"  [실행] v12 — {bridge} (40문항)")
        print(f"{'='*60}")
        t0 = time.time()
        run_rag_v12_full_pipeline(bridge, Q_IDS)
        elapsed = time.time() - t0
        print(f"  ⏱️ 소요: {int(elapsed // 60)}분 {int(elapsed % 60):02d}초")

    # Phase 4: 송정교 채점
    judged_file = results_dir / f"{bridge}_rag_v12_judged.json"
    judged_existing = _load_json(judged_file, default=[]) or []
    if len(judged_existing) >= 40:
        print(f"\n⏭️ [v12] {bridge} 이미 {len(judged_existing)}문항 채점 완료 → 스킵")
    else:
        print(f"\n{'='*60}")
        print(f"  [채점] v12 — {bridge}")
        print(f"{'='*60}")
        judge_all_results(bridge, str(result_file), "rag_v12")
        time.sleep(5)

    # Phase 5: 전체 비교 데이터 수집
    print(f"\n{'='*60}")
    print("  [분석] v1, v5, v10r, v11, v12, LLM, Proposed_improved 비교")
    print(f"{'='*60}")
    all_data: dict[str, dict] = {}
    for bridge_name in BRIDGES:
        all_data[bridge_name] = {}
        for _, system_name in ALL_SYSTEMS.items():
            all_data[bridge_name][system_name] = load_scores_map(bridge_name, system_name)

    retrieval_analysis = {b: analyze_v12_retrieval(b) for b in BRIDGES}

    # Phase 6: 터미널 요약
    print("\n" + "=" * 72)
    print("  전체 성능 요약")
    print("=" * 72)
    for bridge_name in BRIDGES:
        print(f"\n■ {bridge_name}")
        bridge_data = all_data[bridge_name]
        header = f"  {'시스템':<20} | {'정확성':>6} | {'완전성':>6} | {'충실성':>6} | {'전체':>6}"
        print(header)
        print("  " + "-" * (len(header) - 2))
        for label, system_name in ALL_SYSTEMS.items():
            score_map = bridge_data.get(system_name, {})
            metric_avgs = {m: avg_scores(score_map, Q_IDS, m) or 0 for m in METRICS}
            overall = overall_avg(score_map, Q_IDS) or 0
            print(
                f"  {label:<20} | {metric_avgs['accuracy']:>6.2f} | "
                f"{metric_avgs['completeness']:>6.2f} | {metric_avgs['faithfulness']:>6.2f} | "
                f"{overall:>6.2f}"
            )

    print("\n" + "=" * 72)
    print("  v12 검색/이미지 특성 요약")
    print("=" * 72)
    for bridge_name in BRIDGES:
        summary = summarize_v12_retrieval(retrieval_analysis[bridge_name])
        chapter_hits = ", ".join(f"{ch}장={cnt}" for ch, cnt in summary["chapter_counts"].items()) or "-"
        print(f"\n■ {bridge_name}")
        print(f"  문항 수            : {summary['n_questions']}")
        if summary["avg_chapters"] is not None:
            print(f"  평균 예측 장 수    : {summary['avg_chapters']:.2f}")
        if summary["avg_pages"] is not None:
            print(f"  평균 검색 페이지 수: {summary['avg_pages']:.2f}")
        if summary["avg_image_pages"] is not None:
            print(f"  평균 이미지 페이지 : {summary['avg_image_pages']:.2f}")
        if summary["avg_tables"] is not None:
            print(f"  평균 표 청크 수    : {summary['avg_tables']:.2f}")
        print(f"  장 히트 분포       : {chapter_hits}")

    # Phase 7: Excel
    xlsx_path = str(results_dir / "final_comparison_v12.xlsx")
    create_excel(all_data, retrieval_analysis, xlsx_path)

    # Phase 8: Markdown
    md_path = str(results_dir / "paper_data_summary_v6.md")
    create_markdown(all_data, retrieval_analysis, md_path)

    elapsed = time.time() - global_start
    print(f"\n{'='*72}")
    print(f"  전체 완료! {int(elapsed // 60)}분 {int(elapsed % 60):02d}초")
    print(f"  Excel: {xlsx_path}")
    print(f"  Markdown: {md_path}")
    print(f"{'='*72}")


if __name__ == "__main__":
    main()
